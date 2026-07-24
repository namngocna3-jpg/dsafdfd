# -*- coding: utf-8 -*-
# Bổ sung 4 gap + KOL cụ thể vào CẢ 3 VER (không tạo file mới)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

H=Font(bold=True,color="FFFFFF",size=11); HFILL=PatternFill("solid",fgColor="1F3864")
SUB=Font(bold=True,color="1F3864",size=11); TITLE=Font(bold=True,size=14,color="1F3864")
GOLD=PatternFill("solid",fgColor="FFF2CC")
WRAP=Alignment(wrap_text=True,vertical="top"); CENTER=Alignment(horizontal="center",vertical="center",wrap_text=True)
thin=Side(style="thin",color="BBBBBB"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
RED=Font(italic=True,color="C00000")
def hstyle(ws,row,ncol):
    for c in range(1,ncol+1):
        x=ws.cell(row=row,column=c); x.font=H; x.fill=HFILL; x.alignment=CENTER; x.border=BORDER
def body(ws,start,ncol):
    for r in ws.iter_rows(min_row=start,max_col=ncol):
        for c in r: c.alignment=WRAP; c.border=BORDER
def setw(ws,w):
    for i,v in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width=v

def add_sheets(F):
    wb=openpyxl.load_workbook(F)

    # GAP 1: Quy đổi Credit -> Video
    if "B1. Quy đổi Credit-Video" not in wb.sheetnames:
        ws=wb.create_sheet("B1. Quy đổi Credit-Video")
        ws["A1"]="BẢNG QUY ĐỔI CREDIT → VIDEO (từ data vé event, để tư vấn KH)"; ws["A1"].font=TITLE; ws.merge_cells("A1:E1")
        hdr=["Mốc credit","Số video ước tính","Suy ra ~credit/video","Nguồn","Ghi chú tư vấn"]
        ws.append([]); ws.append(hdr); hstyle(ws,3,5)
        rows=[
         ["2.000 (vé Silver)","~200 video Seedance","~10 credit/video","Data vé Silver event","Video cơ bản, độ dài ngắn"],
         ["4.000 (gói Gold web)","~400 video","~10 credit/video","Suy từ Silver","Gấp đôi Silver"],
         ["10.500 (gói Diamond web)","~950-1.050 video","~10 credit/video","Suy tuyến tính","Cho agency dùng nhiều"],
         ["11.500 (vé Gold event)","~1.150 video + face thật","~10 credit/video","Data vé Gold event","Có eKYC face + 4K"],
         ["22.000 (gói Titan web)","~2.200 video","~10 credit/video","Suy tuyến tính","Enterprise/dự án lớn"],
        ]
        for r in rows: ws.append(r)
        body(ws,4,5); setw(ws,[22,22,20,22,26])
        for r in range(4,4+len(rows)): ws.row_dimensions[r].height=34
        n=ws.max_row+2
        ws.cell(row=n,column=1,value="LƯU Ý: ~10 credit/video là mức CƠ BẢN (video ngắn, độ phân giải thường). Video dài hơn / 4K / face thật tốn nhiều credit hơn — cần bảng chi tiết từ Coco. Dùng con số này khi tư vấn để KH hình dung 'tiền ra bao nhiêu video'.").font=RED
        ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=5)

    # GAP 2: Map gói Web vs Vé Event + chính sách credit
    if "B2. Map gói Web vs Vé" not in wb.sheetnames:
        ws=wb.create_sheet("B2. Map gói Web vs Vé")
        ws["A1"]="ĐỐI CHIẾU GÓI WEB vs VÉ EVENT — GIỮ NHẤT QUÁN GIÁ"; ws["A1"].font=TITLE; ws.merge_cells("A1:E1")
        hdr=["Hạng","Vé Event (kèm gì)","Gói Web bán lẻ","Chênh lệch/định vị","Xử lý khi KH hỏi"]
        ws.append([]); ws.append(hdr); hstyle(ws,3,5)
        rows=[
         ["Entry","490k (chỉ vé + teabreak)","-","Vé thuần, không credit","'Vé Entry để trải nghiệm sự kiện'"],
         ["Silver","990k = vé + 2.000cr (~200 video)","990k = 2.000cr","BẰNG GIÁ → nhất quán","'Đúng bằng gói Silver trên web'"],
         ["Gold (event)","4.990k = vé + 11.500cr + face thật+eKYC","-","Vé Gold = ưu đãi lớn (nhiều credit hơn Diamond web)","'Vé Gold event lời hơn mua lẻ'"],
         ["Diamond (web)","-","4.990k = 10.500cr","Cùng giá vé Gold nhưng ít credit hơn (vì không kèm sự kiện)","'Gói web thường, không phải ưu đãi event'"],
         ["Titan (web)","-","9.990k = 22.000cr + face unlock","Cao nhất, cho dự án lớn","'Cho brand/enterprise dùng face thật'"],
        ]
        for r in rows: ws.append(r)
        body(ws,4,5); setw(ws,[14,32,24,30,30])
        for r in range(4,4+len(rows)): ws.row_dimensions[r].height=44
        n=ws.max_row+2
        ws.cell(row=n,column=1,value="INSIGHT SALE: Vé Gold event (4.99tr = 11.500cr) 'lời' hơn gói Diamond web (4.99tr = 10.500cr) vì kèm sự kiện + face thật. Dùng điểm này để đẩy vé event, và sau event upsell người mua Silver lên gói web cao hơn. Chính sách credit: đọc kỹ tinyurl.com/HDSD-Credits trước khi tư vấn.").font=RED
        ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=5)

    # GAP 3: Persona đủ 5 tệp (bổ sung Game Studio + AI Academy)
    if "B3. Persona bổ sung (5 tệp)" not in wb.sheetnames:
        ws=wb.create_sheet("B3. Persona bổ sung (5 tệp)")
        ws["A1"]="PERSONA ĐỦ 5 TỆP BRIEF — BỔ SUNG GAME STUDIO & AI ACADEMY"; ws["A1"].font=TITLE; ws.merge_cells("A1:D1")
        hdr=["Yếu tố","Persona 4: GAME STUDIO","Persona 5: AI ACADEMY / ĐÀO TẠO","(P1-3 xem sheet Customer Persona)"]
        ws.append([]); ws.append(hdr); hstyle(ws,3,4)
        rows=[
         ["Chân dung","Art Lead/Producer studio game indie-mid, 26-40t","Chủ trung tâm/giảng viên đào tạo AI, 28-45t",""],
         ["Mục tiêu","Concept art, trailer, cutscene, asset nhanh & rẻ","Có tài liệu/case dạy học viên, nâng uy tín",""],
         ["Nỗi đau","Ngân sách art hạn chế, cần nhiều asset, deadline build","Nội dung dạy nhanh lỗi thời, thiếu case thực chiến VN",""],
         ["Thông điệp trúng","'Concept & trailer game bằng AI — tiết kiệm 80% chi phí art'","'Nền tảng + cộng đồng để dạy AI Video chuẩn thương mại'",""],
         ["Gói phù hợp","Diamond (dùng nhiều)","Enterprise/đối tác đào tạo",""],
         ["Kênh tiếp cận","Fanpage case game + cộng đồng game dev","Event + hợp tác đào tạo + affiliate",""],
        ]
        for r in rows: ws.append(r)
        body(ws,4,4); setw(ws,[18,36,36,28])
        for r in range(4,4+len(rows)): ws.row_dimensions[r].height=44

    # GAP 4: Content plan event bám agenda thật
    if "B4. Content Event theo Agenda" not in wb.sheetnames:
        ws=wb.create_sheet("B4. Content Event theo Agenda")
        ws["A1"]="CONTENT PLAN EVENT — BÁM AGENDA THẬT 31/07"; ws["A1"].font=TITLE; ws.merge_cells("A1:D1")
        hdr=["Mốc agenda","Nội dung khai thác","Bài đăng Fanpage/Group","Thời điểm đăng"]
        ws.append([]); ws.append(hdr); hstyle(ws,3,4)
        rows=[
         ["Teaser speaker Nguyễn Cao Tùng","NSX 'Tôi thấy hoa vàng trên cỏ xanh' nói về AI","'Nhà sản xuất phim đình đám sẽ nói gì về AI Video?'","Trước event 5-7 ngày"],
         ["Teaser Pinzhang Chen (BytePlus)","Regional Sale Director BytePlus về Seedance","'Người của BytePlus tiết lộ gì về Dreamina Seedance?'","Trước 4-5 ngày"],
         ["Teaser Sunil Thomas","Đạo diễn QC quốc tế showcase quy trình thật","'Xem đạo diễn quốc tế ứng dụng AI vào quảng cáo'","Trước 3 ngày"],
         ["Phiên pháp lý (Ng.T.Tuyền-Coco)","Tuân thủ bản quyền/hóa đơn AI Video tại VN","'AI Video làm sao cho ĐÚNG LUẬT tại VN?'","Trước 2 ngày + nhấn USP"],
         ["Showcase game","Ứng dụng AI Video quảng cáo game","'AI Video trong game — case thực tế'","Trong event (nếu được phép)"],
         ["Panel Hype→Thực tế","Từ trào lưu đến ứng dụng thật","Recap quote hay nhất → Group","Sau event 1-2 ngày"],
         ["Công bố cộng đồng 2026-2027","Hoạt động hỗ trợ cộng đồng AI Việt","'Chính thức: cộng đồng AI Video Việt 2026-2027'","Bế mạc → chốt kéo group"],
        ]
        for r in rows: ws.append(r)
        body(ws,4,4); setw(ws,[28,30,34,20])
        for r in range(4,4+len(rows)): ws.row_dimensions[r].height=40
        n=ws.max_row+2
        ws.cell(row=n,column=1,value="LƯU Ý CHÍNH SÁCH BTC: sự kiện ĐÓNG (closed-door), KHÔNG tự ý quay/livestream showcase độc quyền & phần BytePlus. Chỉ khai thác phần được phép. Trang phục Business Casual. Vé không chuyển nhượng.").font=RED
        ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=4)

    # GAP 5: KOL cụ thể (có tên thật)
    if "B5. KOL cụ thể cần mời" not in wb.sheetnames:
        ws=wb.create_sheet("B5. KOL cụ thể cần mời")
        ws["A1"]="DANH SÁCH KOL/CREATOR CỤ THỂ CẦN TIẾP CẬN (vibe TaiOn)"; ws["A1"].font=TITLE; ws.merge_cells("A1:E1")
        hdr=["Tên/Kênh","Nền tảng & tệp","Vì sao mời","Vai trò đề xuất","Cách tiếp cận"]
        ws.append([]); ws.append(hdr); hstyle(ws,3,5)
        rows=[
         ["TaiOn | AI Worker (@taion_aiwork)","IG 62k, TikTok 73k — AI Creative Educator, có AI Worker Community + AI Academy (Saigon)","Đúng vibe sếp nêu; tệp trẻ sáng tạo, dạy AI workflow thực chiến","Chuyên gia 'Khám video AI' + đồng tổ chức + affiliate","DM nhẹ nhàng (đã có script sheet KOL DM), khen nội dung cụ thể"],
         ["Nguyễn Cao Tùng (NSX)","Đã là speaker event; NSX phim lớn","Uy tín ngành phim, tăng chuẩn cộng đồng","Cố vấn/speaker workshop","Qua BTC event, trang trọng"],
         ["Cộng đồng AI Worker Community","Community sẵn có của TaiOn","Booking chéo, mở tệp nhanh","Đối tác cộng đồng","Deal group booking"],
         ["Các AI Academy/trung tâm đào tạo","Học viện dạy AI (tệp brief nêu)","Nguồn học viên + đối tác đào tạo","Đối tác nội dung + affiliate","Hợp tác win-win"],
         ["Motion/VFX artist studio hậu kỳ","Kênh cá nhân artist VN","Chiều sâu kỹ thuật, giám khảo","Reviewer 'Khám video AI'","DM cá nhân, credit + spotlight"],
         ["Game creator/indie studio VN","Kênh làm game/asset","Đa dạng use-case game","Contest sponsor + case","Deal case study"],
        ]
        for r in rows: ws.append(r)
        body(ws,4,5); setw(ws,[26,32,28,26,26])
        for r in range(4,4+len(rows)): ws.row_dimensions[r].height=54
        n=ws.max_row+2
        ws.cell(row=n,column=1,value="NGUYÊN TẮC: TaiOn là VÍ DỤ vibe sếp muốn — tiếp cận đúng cách (xem KOL DM Scripts). Danh sách này mở rộng dần; xác nhận quan hệ/độ phù hợp trước khi chốt booking. Ưu tiên KOL có COMMUNITY sẵn để booking chéo.").font=RED
        ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=5)

    wb.save(F)
    return len(wb.sheetnames)

for f in ['VER1_OrganicFirst/VER1_CocoPlan.xlsx','VER2_PaidGrowth/VER2_CocoPlan.xlsx','VER3_EventLed/VER3_CocoPlan.xlsx']:
    F="E:/CocoStudio_Plan/"+f
    print(f.split('/')[-1],'->',add_sheets(F),'sheet')
