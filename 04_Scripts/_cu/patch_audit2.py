# -*- coding: utf-8 -*-
# Patch 3 điểm từ audit vòng 2 vào 2 sheet vừa tạo (cả 3 VER)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
SUB=Font(bold=True,color="1F3864",size=11); RED=Font(italic=True,color="C00000")
thin=Side(style="thin",color="BBBBBB"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
WRAP=Alignment(wrap_text=True,vertical="top")

def patch(F):
    wb=openpyxl.load_workbook(F)
    # C1: thêm ghi chú tỉ lệ linh hoạt + quy trình chắp bút
    ws=wb["C1. Content Pillar + Ai làm"]
    n=ws.max_row+2
    adds=[
     "★ QUY TRÌNH NỘI DUNG CHUYÊN GIA (tránh chờ chuyên gia tự viết): Nhân sự phỏng vấn chuyên gia 30' → Nhân sự chắp bút → Chuyên gia duyệt → Đăng. 80% chuyên gia sẽ trễ nếu bắt tự viết.",
     "★ TỈ LỆ PHỄU LINH HOẠT THEO GIAI ĐOẠN: Giai đoạn 0→3k member: tạm ĐẢO nặng LẠNH (reach) để có người trước (Lạnh 45% - Ấm 35% - Nóng 17% - Nền 3%). Khi >3k member mới chuyển về tỉ lệ chuẩn (Ấm 52%). Uy tín cần có người mới xây được.",
     "★ ƯU TIÊN KHI THIẾU NGƯỜI: hy sinh SEEDING trước, GIỮ chuyên gia. 1 buổi Khám video AI = ~50 bài seeding về mặt uy tín (đòn bẩy cao nhất).",
    ]
    for i,t in enumerate(adds):
        c=ws.cell(row=n+i,column=1,value=t); c.font=RED; c.alignment=WRAP
        ws.merge_cells(start_row=n+i,start_column=1,end_row=n+i,end_column=8)
        ws.row_dimensions[n+i].height=44

    # C2: thêm quyền lợi phi vật chất
    ws=wb["C2. Quyền lợi KOL-Chuyên gia"]
    n=ws.max_row+2
    adds2=[
     "★ QUYỀN LỢI PHI VẬT CHẤT (giữ KOL lớn lâu hơn tiền): Ghế HỘI ĐỒNG CỐ VẤN CỘNG ĐỒNG — được tham gia quyết định hướng đi, chọn chủ đề, đề cử member xuất sắc. Status này KOL đầu ngành thèm hơn credit. Chỉ dành Bậc 3.",
     "★ Bổ sung Bậc 3: quyền 'đặt tên' cho chương trình họ dẫn (VD 'Khám video AI cùng [KOL]') = đồng thương hiệu, tạo lòng tự hào & cam kết dài hạn.",
    ]
    for i,t in enumerate(adds2):
        c=ws.cell(row=n+i,column=1,value=t); c.font=RED; c.alignment=WRAP
        ws.merge_cells(start_row=n+i,start_column=1,end_row=n+i,end_column=4)
        ws.row_dimensions[n+i].height=44
    wb.save(F); return True

for f in ['VER1_OrganicFirst/VER1_CocoPlan.xlsx','VER2_PaidGrowth/VER2_CocoPlan.xlsx','VER3_EventLed/VER3_CocoPlan.xlsx']:
    patch("E:/CocoStudio_Plan/"+f); print(f.split('/')[-1],"patched")
