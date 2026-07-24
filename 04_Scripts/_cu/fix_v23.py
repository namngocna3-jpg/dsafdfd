# -*- coding: utf-8 -*-
# VER2/3: thêm ghi chú giai đoạn đo KPI vào sheet '2. KPI chuẩn cộng đồng' (nếu chưa có)
import openpyxl
from openpyxl.styles import Font, Alignment
RED=Font(italic=True,color="C00000"); WRAP=Alignment(wrap_text=True,vertical="top")
MSG="★ GIAI ĐOẠN ĐO KPI (khớp ★AUDIT): 30 ngày đầu (0→~1k member) CHỈ đo 3 số: member mới / % người vào có đăng bài / số câu hỏi thật. Full 10 chỉ số reach-tương tác chỉ bật khi >2-3k member. Lộ trình member chi tiết: xem sheet Lộ trình riêng của VER này."
for f in ['VER2_PaidGrowth/VER2_CocoPlan.xlsx','VER3_EventLed/VER3_CocoPlan.xlsx']:
    F="E:/CocoStudio_Plan/"+f
    wb=openpyxl.load_workbook(F)
    ws=wb["2. KPI chuẩn cộng đồng"]
    # check chưa có
    exists=any(isinstance(c.value,str) and "GIAI ĐOẠN ĐO KPI" in c.value for row in ws.iter_rows() for c in row)
    if not exists:
        n=ws.max_row+2
        c=ws.cell(row=n,column=1,value=MSG); c.font=RED; c.alignment=WRAP
        ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=5); ws.row_dimensions[n].height=44
        wb.save(F); print(f.split('/')[0],"added KPI-phase note")
    else:
        print(f.split('/')[0],"already has note")
