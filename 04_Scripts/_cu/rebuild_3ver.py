# -*- coding: utf-8 -*-
# Gộp lại thành đúng 3 VER tự đủ, BỎ SẠCH tên JOBVUI. Không viết lại nội dung (copy từ file có sẵn).
import openpyxl, copy, os
from openpyxl.utils import get_column_letter

BASE="E:/CocoStudio_Plan"
FULL=openpyxl.load_workbook(f"{BASE}/VER1_OrganicFirst/VER1_CocoPlan_Full.xlsx")
PLUS=openpyxl.load_workbook(f"{BASE}/VER1_OrganicFirst/VER1_PLUS_KhungJOBVUI.xlsx")
V2=openpyxl.load_workbook(f"{BASE}/VER2_PaidGrowth/VER2_CocoPlan_Paid.xlsx")
V3=openpyxl.load_workbook(f"{BASE}/VER3_EventLed/VER3_CocoPlan_Event.xlsx")

REP=[('(khung JOBVUI)',''),('khung JOBVUI','khung chuẩn cộng đồng'),
 ('KHUNG CHUẨN JOBVUI','KHUNG CHUẨN CỘNG ĐỒNG'),('BENCHMARK JOBVUI','THAM CHIẾU NGÀNH'),
 ('BENCHMARK cộng đồng mẫu','THAM CHIẾU NGÀNH'),('vs JOBVUI','vs cộng đồng mẫu thành công'),
 ('JOBVUI','cộng đồng mẫu'),('552k member','cộng đồng quy mô lớn'),('552k','quy mô lớn'),
 ('490tr/6th',''),('490tr',''),('(552k','('),('552.349','—')]

def clean_wb(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str):
                    v=c.value
                    for a,b in REP: v=v.replace(a,b)
                    if v!=c.value: c.value=v

def copy_sheet(src,dst,title):
    ws=dst.create_sheet(title[:31])
    for row in src.iter_rows():
        for c in row:
            nc=ws.cell(row=c.row,column=c.column,value=c.value)
            if c.has_style:
                nc.font=copy.copy(c.font); nc.fill=copy.copy(c.fill)
                nc.alignment=copy.copy(c.alignment); nc.border=copy.copy(c.border)
    for mc in list(src.merged_cells.ranges): ws.merge_cells(str(mc))
    for k,d in src.column_dimensions.items():
        if d.width: ws.column_dimensions[k].width=d.width
    for k,d in src.row_dimensions.items():
        if d.height: ws.row_dimensions[k].height=d.height
    return ws

def build(order, out):
    wb=openpyxl.Workbook(); wb.remove(wb.active)
    for src,name in order:
        if name in src.sheetnames: copy_sheet(src[name],wb,name)
    clean_wb(wb)
    wb.save(out); return len(wb.sheetnames)

# --- danh sách sheet ---
PLUS_STRAT=['★ AUDIT 0-1 (đọc trước)','0. Concept','1. Master Framework','2. KPI chuẩn cộng đồng',
 '3. 6 Kênh tăng member','4. Checklist SEO Group','5. Doanh thu cộng đồng','6. Đối thủ theo tuyến',
 '7. Nhân sự có KPI','8. Quy trình ra ý tưởng']
PLUS_CORE=['★ AUDIT 0-1 (đọc trước)','0. Concept','2. KPI chuẩn cộng đồng','3. 6 Kênh tăng member',
 '4. Checklist SEO Group','5. Doanh thu cộng đồng','6. Đối thủ theo tuyến','7. Nhân sự có KPI','8. Quy trình ra ý tưởng']
FULL_CORE=['2. 10 bài seed (full)','3. Content Calendar 4 tuần','4. Kịch bản Bot','5. Thể lệ Contest #1',
 '6. Benchmark đối thủ','7. Combo giá Coco','8. Kịch bản chốt sale','9. Event Onsite 31-07',
 '12. Fanpage Strategy','13. Customer Persona','14. KOL DM Scripts','15. Research Trend SOP',
 '16. Landing Page Copy','17. Xử lý từ chối Sale','19. Seeding Plan','20. Brand Identity',
 '21. User Journey chi tiết','23. Risk & Crisis','24. Partnership Deck','25. Metrics công thức']
FULL_ALL=[s for s in FULL.sheetnames if s!='00. MỤC LỤC']

# VER1 = PLUS(strategy) + FULL(all) — bản đầy đủ nhất
o1=[(PLUS,n) for n in PLUS_STRAT]+[(FULL,n) for n in FULL_ALL]
# VER2 = Audit+Concept + V2 strategy + core
o2=[(PLUS,'★ AUDIT 0-1 (đọc trước)'),(PLUS,'0. Concept')]+[(V2,n) for n in V2.sheetnames]+\
   [(PLUS,n) for n in PLUS_CORE if n not in('★ AUDIT 0-1 (đọc trước)','0. Concept')]+[(FULL,n) for n in FULL_CORE]
# VER3 = Audit+Concept + V3 strategy + core
o3=[(PLUS,'★ AUDIT 0-1 (đọc trước)'),(PLUS,'0. Concept')]+[(V3,n) for n in V3.sheetnames]+\
   [(PLUS,n) for n in PLUS_CORE if n not in('★ AUDIT 0-1 (đọc trước)','0. Concept')]+[(FULL,n) for n in FULL_CORE]

print("VER1:",build(o1,f"{BASE}/VER1_OrganicFirst/VER1_CocoPlan.xlsx"),"sheet")
print("VER2:",build(o2,f"{BASE}/VER2_PaidGrowth/VER2_CocoPlan.xlsx"),"sheet")
print("VER3:",build(o3,f"{BASE}/VER3_EventLed/VER3_CocoPlan.xlsx"),"sheet")

# xóa file cũ
for f in [f"{BASE}/VER1_OrganicFirst/VER1_CocoPlan_Full.xlsx",
          f"{BASE}/VER1_OrganicFirst/VER1_PLUS_KhungJOBVUI.xlsx",
          f"{BASE}/VER2_PaidGrowth/VER2_CocoPlan_Paid.xlsx",
          f"{BASE}/VER3_EventLed/VER3_CocoPlan_Event.xlsx"]:
    if os.path.exists(f): os.remove(f)
print("Đã xóa file cũ.")

# verify không còn JOBVUI
for f in ['VER1_OrganicFirst/VER1_CocoPlan.xlsx','VER2_PaidGrowth/VER2_CocoPlan.xlsx','VER3_EventLed/VER3_CocoPlan.xlsx']:
    wb=openpyxl.load_workbook(f"{BASE}/{f}")
    hit=0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str) and ('JOBVUI' in c.value or 'Jobvui' in c.value): hit+=1
    print(f, "-> còn JOBVUI:", hit, "| sheet:", len(wb.sheetnames))
