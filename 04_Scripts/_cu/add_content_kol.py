# -*- coding: utf-8 -*-
# Thêm 2 sheet: Content Direction+Pillar (phân rõ ai làm) + Quyền lợi KOL chi tiết. Vào cả 3 VER.
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

H=Font(bold=True,color="FFFFFF",size=11); HFILL=PatternFill("solid",fgColor="1F3864")
SUB=Font(bold=True,color="1F3864",size=11); TITLE=Font(bold=True,size=14,color="1F3864")
GOLD=PatternFill("solid",fgColor="FFF2CC")
BLUE=PatternFill("solid",fgColor="DDEBF7"); ORG=PatternFill("solid",fgColor="FCE4D6"); RED2=PatternFill("solid",fgColor="FCE4E4")
WRAP=Alignment(wrap_text=True,vertical="top"); CENTER=Alignment(horizontal="center",vertical="center",wrap_text=True)
thin=Side(style="thin",color="BBBBBB"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
RED=Font(italic=True,color="C00000")
def hstyle(ws,row,ncol):
    for c in range(1,ncol+1):
        x=ws.cell(row=row,column=c); x.font=H; x.fill=HFILL; x.alignment=CENTER; x.border=BORDER
def body(ws,start,ncol):
    for r in ws.iter_rows(min_row=start,max_col=ncol):
        for c in r: c.alignment=WRAP; c.border=BORDER
def setw(ws,w):
    for i,v in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width=v

def add(F):
    wb=openpyxl.load_workbook(F)

    # ===== C1: CONTENT DIRECTION + PILLAR (phân rõ AI LÀM) =====
    if "C1. Content Pillar + Ai làm" not in wb.sheetnames:
        ws=wb.create_sheet("C1. Content Pillar + Ai làm")
        ws["A1"]="CONTENT DIRECTION + PILLAR — PHÂN RÕ AI SẢN XUẤT"; ws["A1"].font=TITLE; ws.merge_cells("A1:H1")
        hdr=["Phễu","Pillar (tuyến ND)","Tỉ lệ","Mục đích","Định dạng","NGƯỜI SẢN XUẤT CHÍNH","Người hỗ trợ","Tần suất"]
        ws.append([]); ws.append(hdr); hstyle(ws,3,8)
        rows=[
         ["LẠNH\n(Reach-Viral)","Xu hướng & tin AI Video mới","10%","Kéo reach, bắt trend","Carousel/Short","Nhân sự (Content/Mod)","Seeding share","3-4 bài/tuần"],
         ["LẠNH","Giải trí nghề, quizz, đố vui, poll","8%","Tương tác mềm, kéo cmt","Poll/Short","Nhân sự (Mod)","Seeding cmt mồi","2-3 bài/tuần"],
         ["LẠNH","Mẹo/tips nhanh dễ áp dụng","7%","Giá trị nhanh, dễ share","Carousel","Nhân sự (Content)","KOC thử+khoe","2 bài/tuần"],
         ["ẤM\n(Uy tín-Educate)","Kiến thức kỹ thuật chuyên sâu (kèm file)","20%","Xây uy tín, khác biệt","Text dài/Video","CHUYÊN GIA + Content chấp bút","Nhân sự biên tập","2 bài/tuần"],
         ["ẤM","Case study thương mại thật (TVC/game/brand)","15%","Chứng minh năng lực","Album/Video","CHUYÊN GIA/KOL + Nhân sự","Member góp case","1-2 bài/tuần"],
         ["ẤM","🎬 Khám video AI (review bài member)","10%","Mồi tương tác #1","Livestream/Video","CHUYÊN GIA/KOL review","Mod điều phối, Member nộp","1 buổi/tuần"],
         ["ẤM","Xây THCN chuyên gia/KOL (chia sẻ nghề)","7%","Uy tín cá nhân → group","Text/Video","KOL/KOC + Chuyên gia","Nhân sự đăng lại","2 bài/tuần"],
         ["NÓNG\n(Activation-Sale)","Contest/minigame nội bộ","8%","Sân chơi, UGC, viral","Post+thể lệ","Nhân sự (CM) tổ chức","KOL giám khảo, Seeding đẩy","Theo đợt"],
         ["NÓNG","Workshop/Talkshow/Event","8%","Chuyển đổi, gắn kết","Live/Event","Nhân sự + CHUYÊN GIA/KOL speaker","Seeding mời","1/tháng"],
         ["NÓNG","Thông báo/điều hướng/ưu đãi","4%","Kéo phễu, sale khéo","Text/ảnh","Nhân sự (CM/Mod)","Seeding lan","Khi cần"],
         ["NỀN","Bài seeding nuôi group (mồi thảo luận)","3%","Giữ nhịp, không trống","Text/Comment","SEEDING (nick thật)","Mod định hướng","2 bài/ngày"],
        ]
        for r in rows: ws.append(r)
        body(ws,4,8); setw(ws,[12,30,7,18,14,26,22,12])
        # màu theo phễu
        for i,r in enumerate(rows):
            rr=4+i; f=BLUE if r[0].startswith("LẠNH") else ORG if r[0].startswith("ẤM") else RED2 if r[0].startswith("NÓNG") else GOLD
            ws.cell(row=rr,column=1).fill=f; ws.cell(row=rr,column=1).font=SUB
            ws.cell(row=rr,column=6).fill=GOLD  # cột người sản xuất nổi bật
        for r in range(4,4+len(rows)): ws.row_dimensions[r].height=46
        n=ws.max_row+2
        note=[
         "QUY TẮC PHÂN VAI SẢN XUẤT:",
         "• CHUYÊN GIA: nội dung cần độ tin cậy/chiều sâu (kỹ thuật, case, review Khám video AI) — họ là 'bảo chứng' chất lượng.",
         "• KOL/KOC: nội dung cần sức lan tỏa + thương hiệu cá nhân (case, chia sẻ nghề, giám khảo) — họ mang tệp + uy tín.",
         "• NHÂN SỰ CÔNG TY (CM/Mod/Content): xương sống — sản xuất đều, biên tập nội dung chuyên gia, tổ chức hoạt động, giữ nhịp.",
         "• SEEDING (nick thật): mồi thảo luận, comment tạo tương tác, share đa nền tảng — làm group 'sống', không để trống.",
         "TỈ LỆ PHỄU: Lạnh 25% - Ấm 52% - Nóng 20% - Nền 3%. Ấm chiếm nhiều nhất vì đây là group HIGH STANDARD (uy tín > viral rẻ).",
        ]
        for i,t in enumerate(note):
            ws.cell(row=n+i,column=1,value=t).font=(SUB if i==0 else RED)
            ws.merge_cells(start_row=n+i,start_column=1,end_row=n+i,end_column=8)

    # ===== C2: QUYỀN LỢI KOL/KOC/CHUYÊN GIA (chi tiết) =====
    if "C2. Quyền lợi KOL-Chuyên gia" not in wb.sheetnames:
        ws=wb.create_sheet("C2. Quyền lợi KOL-Chuyên gia")
        ws["A1"]="CHÍNH SÁCH QUYỀN LỢI KOL / KOC / CHUYÊN GIA (đi deal)"; ws["A1"].font=TITLE; ws.merge_cells("A1:D1")
        ws.append([]); ws.append(["A. 3 BẬC QUYỀN LỢI (thương lượng linh hoạt)"]); ws.cell(row=3,column=1).font=SUB
        ws.append(["Bậc","Quyền lợi","Đổi lại KOL làm gì","Phù hợp ai"]); hstyle(ws,4,4)
        t1=[
         ["BẬC 1 — Nhẹ","• Tặng credit nền tảng Coco (làm nội dung)\n• Spotlight thương hiệu cá nhân trên Fanpage/Group\n• Badge 'Chuyên gia đồng hành'\n• Được nhắc tên trong recap/tài liệu","Chia sẻ 1-2 bài/tháng đạt tương tác cam kết; review 1 số Khám video AI","KOC nhỏ, chuyên gia mới, artist"],
         ["BẬC 2 — Vừa","• Toàn bộ Bậc 1\n• Đồng sản xuất video demo (2 bên cùng dùng làm portfolio)\n• Ghế speaker workshop/talkshow (PR chéo)\n• Ưu tiên kết nối networking với brand trong cộng đồng\n• Credit lớn hơn + suất dùng tính năng cao (face/4K)","Cam kết nội dung định kỳ; làm giám khảo contest; đồng tổ chức 1 chuyên mục","KOL tầm trung, chuyên gia có tệp"],
         ["BẬC 3 — Nặng","• Toàn bộ Bậc 2\n• AFFILIATE / chia sẻ doanh thu khi giới thiệu khách DN mua gói Coco\n• Đối tác chiến lược dài hạn (hợp đồng)\n• Đồng thương hiệu chương trình (VD 'Khám video AI cùng [KOL]')\n• Ưu tiên booking trả phí cho nội dung lớn","Đại sứ cộng đồng; kéo tệp đều; cam kết dài hạn; giới thiệu khách","KOL lớn, chuyên gia đầu ngành, đối tác"],
        ]
        for r in t1: ws.append(r)
        body(ws,5,4); setw(ws,[16,58,32,24])
        for r in range(5,5+len(t1)): ws.row_dimensions[r].height=110

        sb=5+len(t1)+2
        ws.cell(row=sb,column=1,value="B. QUYỀN LỢI THEO TỪNG LOẠI ĐỐI TÁC").font=SUB
        ws.merge_cells(start_row=sb,start_column=1,end_row=sb,end_column=4)
        ws.append([]); ws.append(["Loại đối tác","Quyền lợi 'trúng' nhất","Hình thức hợp tác","KPI đổi lại"]); hstyle(ws,sb+1,4)
        t2=[
         ["Chuyên gia kỹ thuật (VFX/Motion/Editor)","Credit + spotlight chuyên môn + tệp khách tiềm năng","Reviewer 'Khám video AI', tác giả bài chuyên sâu","2-4 bài/tháng + 1 buổi review"],
         ["KOL Creator (vibe TaiOn)","Đồng sản xuất + affiliate + đồng thương hiệu chương trình","Kéo tệp, dạy workflow, booking chéo community","Share + 1 chương trình/tháng"],
         ["KOC nhỏ (micro)","Credit + quà + được lên spotlight","Thử sản phẩm, khoe kết quả, review thật","2-3 post trải nghiệm/tháng"],
         ["Đạo diễn/NSX cấp cao","Ghế speaker + uy tín + PR thương hiệu cá nhân","Cố vấn, speaker workshop, bảo chứng chuẩn","1-2 lần xuất hiện/quý"],
         ["Agency/Studio đối tác","PR case study 2 chiều + networking brand","Chia sẻ case, đồng tổ chức, giới thiệu khách","1 case + kết nối"],
         ["AI Academy/đào tạo","Đối tác nội dung + affiliate học viên","Cung nội dung dạy, dẫn học viên vào cộng đồng","Nguồn học viên đều"],
        ]
        for r in t2: ws.append(r)
        body(ws,sb+2,4); setw(ws,[26,34,32,24])
        for r in range(sb+2,sb+2+len(t2)): ws.row_dimensions[r].height=52
        n=ws.max_row+2
        ws.cell(row=n,column=1,value="NGUYÊN TẮC DEAL: bắt đầu Bậc 1 để thử hợp tác, nâng bậc theo hiệu quả thực tế. Ưu tiên đối tác có COMMUNITY sẵn (booking chéo). Affiliate là đòn bẩy mạnh nhất — biến KOL thành kênh sale, chỉ trả khi ra khách. Mọi cam kết ghi rõ KPI + thời gian trong thỏa thuận nhẹ (xem Partnership Deck).").font=RED
        ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=4)

    wb.save(F); return len(wb.sheetnames)

for f in ['VER1_OrganicFirst/VER1_CocoPlan.xlsx','VER2_PaidGrowth/VER2_CocoPlan.xlsx','VER3_EventLed/VER3_CocoPlan.xlsx']:
    print(f.split('/')[-1],'->',add("E:/CocoStudio_Plan/"+f),'sheet')
