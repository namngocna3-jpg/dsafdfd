# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("KE HOACH COCO - CHI TIET (pitch) v2.xlsx")
H=Font(bold=True,color="FFFFFF",size=11); HFILL=PatternFill("solid",fgColor="1F3864")
SUB=Font(bold=True,color="1F3864",size=11); TITLE=Font(bold=True,size=14,color="1F3864")
GOLD=PatternFill("solid",fgColor="FFF2CC")
WRAP=Alignment(wrap_text=True,vertical="top"); CENTER=Alignment(horizontal="center",vertical="center",wrap_text=True)
thin=Side(style="thin",color="BBBBBB"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
RED=Font(italic=True,color="C00000")
def hstyle(ws,row,ncol):
    for c in range(1,ncol+1):
        x=ws.cell(row=row,column=c); x.font=H; x.fill=HFILL; x.alignment=CENTER; x.border=BORDER
def body(ws,start,ncol):
    for r in ws.iter_rows(min_row=start,max_col=ncol):
        for c in r: c.alignment=WRAP; c.border=BORDER
def setw(ws,w):
    for i,v in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width=v

# ====== SHEET 12: FANPAGE STRATEGY ======
ws=wb.create_sheet("12. Fanpage Strategy")
ws["A1"]="FANPAGE COCO STUDIO — CHÍNH CHUYÊN (tách khỏi Group)"; ws["A1"].font=TITLE; ws.merge_cells("A1:D1")
ws.append([]); ws.append(["A. ĐỊNH VỊ & PILLAR"]); ws.cell(row=3,column=1).font=SUB
h1=["Pillar","% Content","Vai trò","Ví dụ chủ đề"]
ws.append(h1); hstyle(ws,4,4)
pil=[
 ["Sản phẩm & Công nghệ (Seedance/Coco)","30%","Khẳng định năng lực","Tính năng mới Seedance, so sánh model, unlock face thật/4K"],
 ["Case study & Showcase","25%","Chứng minh thực chiến","TVC/trailer làm bằng Coco, before-after, số liệu tiết kiệm"],
 ["Kiến thức chuyên môn","20%","Xây uy tín","Prompt technique, camera language, giữ continuity"],
 ["Xu hướng & Tin ngành","15%","Bắt trend, dẫn dắt","Trend AI video toàn cầu, cập nhật công nghệ, góc nhìn đạo diễn"],
 ["Thương hiệu & Sự kiện","10%","Kết nối cộng đồng","Event 31/07, hợp tác BytePlus, hoạt động cộng đồng 2026-2027"],
]
for r in pil: ws.append(r)
body(ws,5,4)
sb=5+len(pil)+2
ws.cell(row=sb,column=1,value="B. 20 IDEA BÀI FANPAGE CỤ THỂ").font=SUB
h2=["#","Pillar","Tiêu đề/Angle","Định dạng"]
ws.append([]); ws.append(h2); hstyle(ws,sb+1,4)
ideas=[
 ["1","Sản phẩm","Seedance 2.0 có gì mới: 5 nâng cấp đáng chú ý","Carousel"],
 ["2","Sản phẩm","Vì sao 'thuần seed' quan trọng khi làm TVC thương mại","Video"],
 ["3","Sản phẩm","Unlock face thật + eKYC: cách Coco xử lý bản quyền khuôn mặt","Text+ảnh"],
 ["4","Sản phẩm","So sánh: cùng 1 prompt trên 4 model, ai thắng?","Video so sánh"],
 ["5","Case study","Làm TVC nước hoa 15s bằng Coco: quy trình A-Z","Album quy trình"],
 ["6","Case study","Trailer game indie dựng bằng AI: tiết kiệm 80% chi phí","Video+số liệu"],
 ["7","Case study","Brand X làm 3 hướng hình trong 1 buổi thay vì 3 tuần","Before-after"],
 ["8","Case study","Từ moodboard đến video: 1 chiến dịch thật","Carousel"],
 ["9","Kiến thức","5 lỗi ánh sáng khiến video AI trông giả","Carousel"],
 ["10","Kiến thức","Ngôn ngữ camera: dolly/orbit/handheld dùng khi nào","Video minh hoạ"],
 ["11","Kiến thức","Cách giữ nhân vật nhất quán qua 10 cảnh","Text+ảnh"],
 ["12","Kiến thức","Checklist brief video thương mại chuẩn client","Carousel tải về"],
 ["13","Xu hướng","3 xu hướng AI video định hình quảng cáo 2026","Carousel"],
 ["14","Xu hướng","Đạo diễn thật nghĩ gì về AI video? (trích event)","Video quote"],
 ["15","Xu hướng","Vì sao brand lớn bắt đầu chuyển ngân sách sang AI video","Text dài"],
 ["16","Xu hướng","Tin ngành: cập nhật công nghệ tuần","Carousel"],
 ["17","Thương hiệu","Coco x BytePlus: vì sao là đối tác chính thức tại VN","Text+ảnh"],
 ["18","Thương hiệu","Hậu trường sự kiện 31/07 tại Sofitel","Album"],
 ["19","Thương hiệu","Hoạt động hỗ trợ cộng đồng sáng tạo AI Việt 2026-2027","Text"],
 ["20","Thương hiệu","Gặp gỡ đội ngũ & chuyên gia đồng hành cùng Coco","Carousel"],
]
for r in ideas: ws.append(r)
body(ws,sb+2,4); setw(ws,[6,20,58,16])
n=ws.max_row+2
ws.cell(row=n,column=1,value="KHÁC BIỆT FANPAGE vs GROUP: Fanpage = mặt tiền thương hiệu (chính chuyên, 1 chiều, bán sản phẩm gián tiếp). Group = sân chơi cộng đồng (2 chiều, tương tác, không bán trực tiếp). Fanpage 1-2 bài/ngày, giờ vàng 12h & 20h.").font=RED
ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=4)

# ====== SHEET 13: CUSTOMER PERSONA ======
ws=wb.create_sheet("13. Customer Persona")
ws["A1"]="3 CHÂN DUNG KHÁCH HÀNG — ĐỂ SALE & CONTENT ĐÁNH ĐÚNG"; ws["A1"].font=TITLE; ws.merge_cells("A1:D1")
hdr=["Yếu tố","Persona 1: AGENCY/PRODUCTION","Persona 2: BRAND/ENTERPRISE","Persona 3: FREELANCER PRO"]
ws.append([]); ws.append(hdr); hstyle(ws,3,4)
per=[
 ["Chân dung","GĐ sáng tạo / Producer, 28-40t, chạy nhiều job client","Marketing Manager/Brand Lead, 30-45t","Editor/Creator tự do, 24-35t"],
 ["Mục tiêu","Giao job nhanh, nhiều hướng hình, biên lợi nhuận tốt","TVC/quảng cáo đúng chuẩn, đúng luật, đúng brand","Nhận nhiều job hơn, làm nhanh, giá cạnh tranh"],
 ["Nỗi đau","Chi phí ê-kíp cao, deadline gấp, client đổi ý liên tục","Sợ rủi ro pháp lý (bản quyền, mặt người), cần hóa đơn","Không đủ ngân sách tool xịn, cạnh tranh giá"],
 ["Rào cản mua","Sợ chất lượng không đủ giao khách","Sợ pháp lý + cần quy trình chuẩn nội bộ","Nhạy giá, cần dùng thử trước"],
 ["Thông điệp trúng","'3-5 hướng hình/buổi, biên lợi nhuận cao hơn'","'VAT + eKYC + bản quyền = an tâm lên sóng'","'Chất studio, giá freelancer, có hóa đơn'"],
 ["Gói phù hợp","Diamond","Titan / Enterprise","Silver / Gold"],
 ["Kênh tiếp cận","Fanpage case study + inbox sale + event","Sale trực tiếp + event + LinkedIn","Group cộng đồng + contest + Fanpage"],
]
for r in per: ws.append(r)
body(ws,4,4); setw(ws,[18,32,32,32])
for r in range(4,4+len(per)): ws.row_dimensions[r].height=52

# ====== SHEET 14: KOL DM SCRIPTS ======
ws=wb.create_sheet("14. KOL DM Scripts")
ws["A1"]="TIN NHẮN MỜI KOL/CHUYÊN GIA — NGUYÊN VĂN THEO NHÓM"; ws["A1"].font=TITLE; ws.merge_cells("A1:C1")
hdr=["Nhóm KOL","Tin nhắn mời (nguyên văn)","Điểm nhấn khi deal"]
ws.append([]); ws.append(hdr); hstyle(ws,3,3)
dm=[
 ["AI Creator/Worker (kiểu TaiOn)",
  "Chào anh/bạn [Tên], mình theo dõi kênh [tên kênh] khá lâu, thấy cách anh/bạn chia sẻ workflow AI rất thực chiến 👏\nMình đang xây cộng đồng AI Video Creators Vietnam (bảo trợ bởi Coco Studio - BytePlus), tệp toàn người làm nghề thật. Muốn mời anh/bạn làm chuyên gia đồng hành cho chuyên mục 'Khám video AI' — mình sẽ tặng credit nền tảng Coco để anh/bạn làm nội dung, đồng sản xuất video demo cùng, và có chính sách affiliate khi giới thiệu khách doanh nghiệp.\nAnh/bạn cho mình xin 15 phút call tuần này để trao đổi cụ thể nhé?",
  "Nhấn: credit + đồng sản xuất + affiliate. Đây là nhóm KÉO TỆP mạnh nhất."],
 ["Đạo diễn/NSX (qua BTC event)",
  "Kính gửi anh/chị [Tên], em là [tên] phụ trách cộng đồng của Coco Studio. Sau sự kiện 31/07, bên em xây dựng cộng đồng dài hạn cho người làm phim & quảng cáo ứng dụng AI tại VN.\nEm rất mong được mời anh/chị tham gia với vai trò cố vấn/speaker cho các buổi chia sẻ — đây cũng là cơ hội lan toả thương hiệu cá nhân của anh/chị tới tệp sáng tạo chuyên nghiệp. Em xin phép gửi đề xuất chi tiết qua email được không ạ?",
  "Nhấn: uy tín + spotlight cá nhân. Trang trọng vì đây là KOL cấp cao."],
 ["VFX/Motion Artist",
  "Hi [Tên], mình xem reel của bạn trên [nền tảng], phần [chi tiết cụ thể] làm quá đỉnh. Bên mình đang chạy chuyên mục 'Khám video AI' trong cộng đồng AI Video Creators Vietnam — muốn mời bạn làm giám khảo/reviewer để soi bài kỹ thuật cho anh em. Đổi lại có credit Coco + spotlight + kết nối tệp khách hàng tiềm năng. Bạn hứng thú không, mình gửi chi tiết nhé?",
  "Nhấn: chuyên môn được công nhận + tệp khách."],
 ["Agency C-level",
  "Chào anh/chị [Tên], em thấy [agency] có nhiều campaign ấn tượng. Coco Studio đang xây cộng đồng cho giới sản xuất chuyên nghiệp ứng dụng AI. Em muốn mời [agency] tham gia với vai trò đối tác nội dung — chia sẻ case study (được PR ngược lại) và networking với các brand trong cộng đồng. Anh/chị sắp xếp được buổi cà phê trao đổi không ạ?",
  "Nhấn: PR 2 chiều + networking brand."],
]
for r in dm: ws.append(r)
body(ws,4,3); setw(ws,[24,74,26])
for r in range(4,4+len(dm)): ws.row_dimensions[r].height=140

# ====== SHEET 15: RESEARCH TREND SOP ======
ws=wb.create_sheet("15. Research Trend SOP")
ws["A1"]="SOP NGHIÊN CỨU XU HƯỚNG — QUY TRÌNH HÀNG TUẦN (nhiệm vụ #5 sếp giao)"; ws["A1"].font=TITLE; ws.merge_cells("A1:D1")
hdr=["Tần suất","Việc làm","Nguồn theo dõi","Output áp dụng"]
ws.append([]); ws.append(hdr); hstyle(ws,3,4)
sop=[
 ["Hàng ngày (15')","Lướt nguồn trend, lưu 2-3 thứ hay","X/Twitter (#AIvideo), Reddit r/aivideo, TikTok, kênh KOL","Note vào 'Trend bank'"],
 ["Thứ 2 (30')","Tổng hợp trend tuần → chọn 3 cái làm được","Trend bank + tin công nghệ (BytePlus, ByteDance blog)","Bài 'Trend Radar' CN"],
 ["Thứ 4 (30')","Thử 1 prompt/kỹ thuật trend → làm demo","Nền tảng Coco/Seedance","Demo cho Fanpage/Group"],
 ["Hàng tuần (30')","Theo dõi đối thủ (SDVN, group AI khác) có gì hot","Group đối thủ, Fanpage cạnh tranh","Điều chỉnh content angle"],
 ["Hàng tháng (1h)","Báo cáo trend tháng + đề xuất tuyến nội dung mới","Tổng hợp cả tháng","Cập nhật Content Direction"],
]
for r in sop: ws.append(r)
body(ws,4,4); setw(ws,[16,32,40,28])
for r in range(4,4+len(sop)): ws.row_dimensions[r].height=48
n=ws.max_row+2
ws.cell(row=n,column=1,value="NGUỒN TREND GỢI Ý: X (@/hashtag AI video), Reddit (r/StableDiffusion, r/aivideo), TikTok trend, YouTube (kênh AI creator), Futurepedia/There's An AI For That, blog BytePlus/ByteDance, group SDVN. Lập 1 Google Sheet 'Trend Bank' cột: Ngày|Trend|Nguồn|Làm được?|Ưu tiên.").font=RED
ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=4)

# ====== SHEET 16: LANDING PAGE COPY ======
ws=wb.create_sheet("16. Landing Page Copy")
ws["A1"]="COPY LANDING PAGE COCO STUDIO — TỪNG SECTION (bổ sung cho 5 video hero)"; ws["A1"].font=TITLE; ws.merge_cells("A1:C1")
hdr=["Section","Copy đề xuất (nguyên văn)","Ghi chú"]
ws.append([]); ws.append(hdr); hstyle(ws,3,3)
lp=[
 ["Hero (headline)","AI Video chuẩn thương mại. Đủ luật để lên sóng tại Việt Nam.\n(Sub) Nền tảng Seedance thuần seed, mở khoá khuôn mặt thật với eKYC, xuất hoá đơn VAT. Vận hành bởi Peacom — đối tác chính thức BytePlus.","CTA: Bắt đầu ngay / Xem demo"],
 ["Vì sao Coco (4 cột)","1. Pure Native Seedance — dùng full tính năng gốc, không lai model.\n2. Đối tác chính thức BytePlus tại VN.\n3. Mở khoá khuôn mặt thật (eKYC) — an toàn bản quyền.\n4. Thanh toán nội địa + hoá đơn VAT + hỗ trợ tiếng Việt.","Giữ 4 card như web hiện tại"],
 ["Full-Scenario (tabs + 5 video)","Một nền tảng cho mọi kịch bản: Branded/TVC · Cinematic · Game · Real Human · Anime.\n(Mỗi tab gắn 1 video hero + overlay đã viết ở brief)","Nối sheet brief 5 video"],
 ["Everything your production needs","Từ ý tưởng đến file giao khách: prompt theo category, giữ nhất quán nhân vật, xuất 4K, review bởi cộng đồng chuyên gia.","Grid category như web"],
 ["Bảng giá","No Subscription — Pay per use. Không ràng buộc thuê bao, không trộn model, có hoá đơn VAT.\n(4 gói Silver/Gold/Diamond/Titan — nối sheet 7)","Highlight Gold/Diamond"],
 ["Social proof","Được tin dùng bởi các production house & agency tại VN. [Logo khách + trích dẫn case]","Bổ sung khi có logo thật"],
 ["FAQ","What is Coco Studio? / 'Native Seedance' nghĩa là gì? / eKYC bảo vệ mình ra sao? / Credit hoạt động thế nào? / Có xuất hoá đơn doanh nghiệp? / Hỗ trợ kỹ thuật thế nào?","Đã có trên web, giữ nguyên"],
 ["CTA cuối","Sẵn sàng đưa AI Video vào job thật? Bắt đầu với Coco Studio hôm nay.","Nút + form/booking"],
]
for r in lp: ws.append(r)
body(ws,4,3); setw(ws,[22,74,24])
for r in range(4,4+len(lp)): ws.row_dimensions[r].height=80

# ====== SHEET 17: OBJECTION HANDLING ======
ws=wb.create_sheet("17. Xử lý từ chối Sale")
ws["A1"]="XỬ LÝ TỪ CHỐI (OBJECTION HANDLING) — BẢNG PHẢN XẠ NHANH"; ws["A1"].font=TITLE; ws.merge_cells("A1:C1")
hdr=["Khách nói","Cách phản hồi (nguyên văn)","Nguyên tắc"]
ws.append([]); ws.append(hdr); hstyle(ws,3,3)
obj=[
 ["'Đắt hơn các tool khác'","'Dạ đúng là giá/video có cao hơn vài tool nước ngoài. Nhưng những tool đó không xuất hoá đơn VAT, không có quy trình eKYC khuôn mặt — nghĩa là mình không dùng thương mại chính thức được. Cái anh/chị trả thêm là để lên sóng an toàn, không dính bản quyền.'","Đổi khung: giá → rủi ro tránh được"],
 ["'Để tôi tự làm bằng tool free'","'Hoàn toàn được cho việc thử nghiệm ạ. Nhưng khi giao khách doanh nghiệp, thứ họ hỏi là hoá đơn và bản quyền — đó là lúc bản free tắc. Bên em lo trọn phần đó để anh/chị tập trung làm sáng tạo.'","Thừa nhận + chỉ ra giới hạn thật"],
 ["'Chất lượng có đủ giao khách không?'","'Em gửi anh/chị 5 video demo đúng ngành của mình + 1 case thật có số liệu. Anh/chị xem rồi đánh giá trực tiếp. Nếu hợp, em cho dùng thử 1 job thật trước khi quyết.'","Show demo + cho dùng thử"],
 ["'Team chưa biết dùng AI'","'Bên em có tài liệu, cộng đồng chuyên gia và chuyên mục Khám video AI để team học nhanh. Mua gói còn kèm suất mentor. Anh/chị không phải tự mò.'","Bán cả hệ sinh thái hỗ trợ"],
 ["'Để tôi suy nghĩ thêm'","'Dạ em hiểu. Anh/chị cứ giữ [credit trial], thử 1 job nhỏ tuần này. Em đồng hành setup. Thấy ra kết quả rồi mình tính tiếp, không vội ạ.'","Hạ rào cản, giữ liên hệ"],
 ["'Ngân sách năm nay hết rồi'","'Không sao ạ. Em note lại, đầu kỳ sau kết nối lại. Trong lúc đó mời anh/chị vào cộng đồng để cập nhật case & công nghệ, khi cần là mình chạy được ngay.'","Nuôi dài hạn qua cộng đồng"],
]
for r in obj: ws.append(r)
body(ws,4,3); setw(ws,[24,74,24])
for r in range(4,4+len(obj)): ws.row_dimensions[r].height=72

wb.save("KE HOACH COCO - CHI TIET (pitch) v2.xlsx")
print("DEEP APPENDED. Total sheets:", len(wb.sheetnames))
for s in wb.sheetnames: print(" -", s)
